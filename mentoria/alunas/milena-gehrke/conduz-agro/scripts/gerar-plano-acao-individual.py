# Gera o Plano de Acao Individual (pos Vendas Secretas) na identidade visual Conduz Agro
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LOGO = sys.argv[1]
OUT = sys.argv[2]

INK = RGBColor(0x3D, 0x28, 0x17)
INK_SOFT = RGBColor(0x6B, 0x5A, 0x44)
OLIVE = RGBColor(0x5C, 0x6B, 0x3F)
GOLD = RGBColor(0xB0, 0x7A, 0x16)
PH = RGBColor(0x9A, 0x8A, 0x70)  # placeholder
SERIF, SANS = "Fraunces", "Quicksand"
RAISED, DEEP, GOLD_BG, OLIVE_BG = "EDE6D6", "E4DBC4", "F3E6C8", "E3E6D5"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.2))
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)

st = doc.styles["Normal"]
st.font.name = SANS
st.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
st.font.size = Pt(10.5)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.25


def run(p, text, font=SANS, size=None, bold=False, italic=False, color=INK, caps=False, spacing=None):
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    r.font.bold, r.font.italic = bold, italic
    r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if caps:
        r.font.all_caps = True
    if spacing is not None:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        r._element.rPr.append(sp)
    return r


def para(container, before=0, after=4, align=None):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    return p


def shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_borders(cell, **edges):
    tcPr = cell._element.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            color, sz = edges[edge]
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcPr.append(b)


def cell_margins(cell, top=160, bottom=160, left=220, right=220):
    tcPr = cell._element.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tcPr.append(m)


def box(fill=RAISED, border=None, cols=1, widths=None):
    t = doc.add_table(rows=1, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = widths or [Cm(16.6 / cols)] * cols
    for i, c in enumerate(t.rows[0].cells):
        c.width = widths[i]
        shade(c, fill)
        cell_borders(c, **(border or {}))
        cell_margins(c)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
    no_split(t)
    return t


def no_split(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit")
        trPr.append(el)


def first(cell):
    p = cell.paragraphs[0]
    return p if not p.text and not p.runs else cell.add_paragraph()


def spacer(pt=6):
    p = para(doc, 0, 0)
    p.paragraph_format.line_spacing = Pt(pt)


def section(num, title):
    p = para(doc, 22, 2)
    p.paragraph_format.keep_with_next = True
    run(p, f"{num:02d}", SERIF, 11, True, color=GOLD)
    p = para(doc, 0, 10)
    p.paragraph_format.keep_with_next = True
    run(p, title, SERIF, 17, True)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "6"), ("color", DEEP)):
        bot.set(qn(f"w:{k}"), v)
    bdr.append(bot)
    pPr.append(bdr)


def label(container, text, color=OLIVE, before=6):
    p = para(container, before, 1)
    p.paragraph_format.keep_with_next = True
    run(p, text, SANS, 8, True, color=color, caps=True, spacing=20)
    return p


def placeholder(container, text="[ ]", after=4):
    p = para(container, 0, after)
    run(p, text, SANS, 10.5, italic=True, color=PH)
    return p


def body(container, text, after=4, color=INK, size=10.5, bold=False, italic=False, font=SANS):
    p = para(container, 0, after)
    run(p, text, font, size, bold, italic, color)
    return p


# ---------- CAPA / CABECALHO ----------
t = doc.add_table(rows=1, cols=2)
t.autofit = False
c0, c1 = t.rows[0].cells
c0.width, c1.width = Cm(2.2), Cm(14.4)
for c in (c0, c1):
    cell_borders(c)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
c0.paragraphs[0].add_run().add_picture(LOGO, width=Cm(1.8))
p = c1.paragraphs[0]
p.paragraph_format.space_after = Pt(0)
run(p, "Conduz ", SERIF, 15, True)
run(p, "Agro", SERIF, 15, True, color=OLIVE)
p = c1.add_paragraph()
run(p, "Por Milena Gehrke", SANS, 9, color=INK_SOFT)

spacer(14)
p = para(doc, 0, 2)
run(p, "Plano de Ação Individual", SERIF, 28, True)
p = para(doc, 0, 14)
run(p, "Seu diagnóstico e seu caminho de evolução profissional", SANS, 12, True, color=OLIVE)

t = box(RAISED, {"left": ("5C6B3F", 18)})
c = t.rows[0].cells[0]
for i, (k, v) in enumerate((("Profissional", "[ ]"), ("Data", "[ ]"), ("Diagnóstico realizado por", "Milena Gehrke | Conduz Agro"))):
    p = c.paragraphs[0] if i == 0 else c.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run(p, f"{k}:  ", SANS, 10, True, color=INK_SOFT)
    run(p, v, SANS, 10.5, italic=(v == "[ ]"), color=PH if v == "[ ]" else INK)

# ---------- 1. PRE-DIAGNOSTICO ----------
section(1, "Seu Pré-Diagnóstico")
body(doc, "A partir da nossa conversa, identificamos os principais pontos que hoje influenciam sua atuação profissional.", 10, INK_SOFT)

for pilar in ("Técnica", "Emocional", "Condução"):
    t = box(RAISED, {"top": ("B07A16", 18)})
    c = t.rows[0].cells[0]
    p = c.paragraphs[0]
    run(p, pilar, SERIF, 13, True)
    label(c, "Como você está hoje")
    placeholder(c, "[Diagnóstico individual]")
    label(c, "Principal ponto a desenvolver")
    placeholder(c)
    label(c, "O que isso pode estar limitando")
    placeholder(c, after=0)
    spacer(8)

# ---------- 2. MOMENTO PROFISSIONAL ----------
section(2, "Seu Momento Profissional")
body(doc, "Embora existam pontos de desenvolvimento nos três pilares, todo processo de evolução precisa começar por uma prioridade.", 10, INK_SOFT)

t = box(GOLD_BG, {"left": ("B07A16", 18)})
c = t.rows[0].cells[0]
p = c.paragraphs[0]
run(p, "Seu principal ponto de atenção hoje é", SANS, 8, True, color=GOLD, caps=True, spacing=20)
p = para(c, 2, 0)
run(p, "[TÉCNICA / EMOCIONAL / CONDUÇÃO]", SERIF, 15, True, color=PH)
spacer(8)

label(doc, "Por quê?", before=4)
placeholder(doc, "[Explicação personalizada do diagnóstico.]", 8)
body(doc, "Esse ponto foi identificado como prioritário porque, enquanto ele não for desenvolvido, tende a continuar impactando:", 4)
for i in (1, 2, 3):
    p = para(doc, 0, 2)
    p.paragraph_format.left_indent = Cm(0.5)
    run(p, "—  ", SANS, 10.5, True, color=GOLD)
    run(p, f"[resultado {i}]", SANS, 10.5, italic=True, color=PH)
spacer(6)
body(doc, "Em outras palavras:", 4, INK_SOFT)

t = box("FFFFFF", {"left": ("5C6B3F", 24)})
c = t.rows[0].cells[0]
cell_margins(c, 120, 120, 300, 200)
p = c.paragraphs[0]
run(p, "Você não precisa simplesmente fazer mais. Precisa desenvolver aquilo que hoje está impedindo sua evolução.", SERIF, 13, True, True, OLIVE)

# ---------- 3. PLANO DE ACAO ----------
section(3, "Seu Plano de Ação e Evolução")

fases = [
    ("Primeiros 4 meses", "Destravar",
     "Trabalhar os principais gargalos identificados e criar uma base mais segura para sua evolução.",
     "[Técnica / Emocional / Condução]",
     "Mais clareza para agir, mais segurança para se posicionar e maior capacidade de lidar com as situações que hoje geram dificuldade."),
    ("Até 8 meses", "Consolidar",
     "Transformar o desenvolvimento inicial em comportamento, processo e resultado profissional.",
     "[Técnica / Emocional / Condução]",
     "Mais autonomia, melhor comunicação, maior capacidade de condução e uma atuação menos reativa e mais estratégica."),
    ("Até 12 meses", "Evoluir para Autoridade",
     "Integrar técnica, emocional e condução para consolidar uma atuação profissional mais estratégica.",
     "Integração dos três pilares.",
     "Um profissional que não apenas executa, mas entende, diagnostica, orienta, se posiciona e conduz."),
]
for prazo, nome, obj, foco, construir in fases:
    t = box(RAISED, {"top": ("5C6B3F", 18)})
    c = t.rows[0].cells[0]
    p = c.paragraphs[0]
    run(p, prazo, SANS, 8, True, color=GOLD, caps=True, spacing=20)
    p = para(c, 0, 4)
    run(p, nome, SERIF, 15, True)
    label(c, "Objetivo")
    body(c, obj)
    label(c, "Foco principal")
    if foco.startswith("["):
        placeholder(c, foco)
    else:
        body(c, foco)
    label(c, "O que desenvolver")
    for _ in range(3):
        placeholder(c, after=1)
    label(c, "Ações prioritárias")
    for _ in range(4):
        placeholder(c, after=1)
    label(c, "O que esperamos construir")
    body(c, construir, 0, OLIVE, italic=True)
    spacer(10)

# ---------- 4. MAPA ----------
section(4, "Seu Mapa de Evolução")
etapas = [
    ("Hoje", "Onde estou?", "[Descrição breve do estágio atual.]"),
    ("4 meses", "O que preciso destravar?", "[ ]"),
    ("8 meses", "O que preciso consolidar?", "[ ]"),
    ("12 meses", "Quem preciso me tornar profissionalmente?", "[ ]"),
]
t = doc.add_table(rows=len(etapas), cols=2)
t.autofit = False
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (quando, pergunta, ph) in enumerate(etapas):
    a, b = t.rows[i].cells
    a.width, b.width = Cm(3.2), Cm(13.4)
    shade(a, "5C6B3F" if i == len(etapas) - 1 else OLIVE_BG)
    shade(b, RAISED)
    for c in (a, b):
        cell_borders(c, bottom=("F5F0E6", 24))
        cell_margins(c, 140, 140, 200, 200)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = a.paragraphs[0]
    run(p, quando, SERIF, 12, True, color=RGBColor(0xF5, 0xF0, 0xE6) if i == len(etapas) - 1 else OLIVE)
    p = b.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run(p, pergunta, SANS, 10.5, True)
    placeholder(b, ph, 0)

no_split(t)

# ---------- 5. PILARES ----------
section(5, "Os Três Pilares da Sua Evolução")
pilares = [
    ("Eu Sei", "Técnica", "Conheço, analiso, diagnostico e entrego."),
    ("Eu Sustento", "Emocional", "Confio no meu conhecimento, sustento minhas decisões, estabeleço limites e lido melhor com pressão e conflitos."),
    ("Eu Conduzo", "Condução", "Sei perguntar, diagnosticar, comunicar, orientar, negociar e conduzir o cliente para o próximo passo."),
]
t = doc.add_table(rows=1, cols=3)
t.autofit = False
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, (lema, pilar, desc) in zip(t.rows[0].cells, pilares):
    c.width = Cm(5.53)
    shade(c, RAISED)
    cell_borders(c, top=("B07A16", 18), left=("F5F0E6", 30), right=("F5F0E6", 30))
    cell_margins(c, 180, 180, 200, 200)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run(p, lema, SERIF, 14, True)
    p = para(c, 0, 6)
    run(p, pilar, SANS, 8, True, color=OLIVE, caps=True, spacing=20)
    body(c, desc, 0, INK_SOFT, 9.5)

# ---------- 6. PRIMEIRA ACAO ----------
section(6, "Sua Primeira Ação")
body(doc, "Antes de pensar nos próximos 12 meses, existe uma ação que precisa começar agora:", 8, INK_SOFT)
t = box(GOLD_BG, {"left": ("B07A16", 18)})
c = t.rows[0].cells[0]
p = c.paragraphs[0]
run(p, "[AÇÃO PRIORITÁRIA]", SERIF, 15, True, color=PH)
label(c, "Prazo", GOLD, 8)
placeholder(c)
label(c, "Como colocar em prática", GOLD)
placeholder(c)
label(c, "Como saberemos que houve evolução", GOLD)
placeholder(c, after=0)

# ---------- 7. PROXIMO NIVEL ----------
section(7, "Uma Visão Sobre Seu Próximo Nível")
body(doc, "Seu desenvolvimento não depende apenas de acumular mais conhecimento técnico.", 6)
body(doc, "À medida que você cresce profissionalmente, muda também a forma como você:", 6)
p = para(doc, 4, 10, WD_ALIGN_PARAGRAPH.CENTER)
passos = ["pensa", "se posiciona", "comunica", "decide", "conduz"]
for i, s in enumerate(passos):
    run(p, s, SERIF, 13, True, color=OLIVE)
    if i < len(passos) - 1:
        run(p, "  →  ", SANS, 12, True, color=GOLD)
body(doc, "O objetivo é sair de uma atuação em que você apenas responde às demandas para uma atuação em que você consegue identificar o problema, orientar o cliente e conduzir o caminho para a solução.", 16)

t = box("3D2817")
c = t.rows[0].cells[0]
cell_margins(c, 360, 360, 300, 300)
p = c.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
cream = RGBColor(0xF5, 0xF0, 0xE6)
run(p, "EU SEI.  ", SERIF, 17, True, color=cream)
run(p, "EU SUSTENTO.  ", SERIF, 17, True, color=RGBColor(0xD4, 0xA0, 0x17))
run(p, "EU CONDUZO.", SERIF, 17, True, color=cream)
p = para(c, 0, 0, WD_ALIGN_PARAGRAPH.CENTER)
run(p, "Esse é o caminho de evolução proposto pelo Conduz Agro.", SANS, 10, color=RGBColor(0xE4, 0xDB, 0xC4))

# ---------- RODAPE ----------
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp, "Conduz Agro  ·  Milena Gehrke", SANS, 8, color=INK_SOFT)

doc.save(OUT)
print("ok", OUT)
